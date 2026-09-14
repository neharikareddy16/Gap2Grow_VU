import os
import re
import json
import difflib
import urllib.parse
import urllib.request
from typing import Dict, Any, List, Optional
from dotenv import load_dotenv
load_dotenv()

from subject_data import SUBJECTS_DATA

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from youtube_transcript_api import YouTubeTranscriptApi
    YOUTUBE_TRANSCRIPT_AVAILABLE = True
except ImportError:
    YOUTUBE_TRANSCRIPT_AVAILABLE = False

try:
    import google.generativeai as genai
    GEMINI_AVAILABLE = True
except ImportError:
    GEMINI_AVAILABLE = False

try:
    from openai import OpenAI
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False

try:
    from groq import Groq
    GROQ_AVAILABLE = True
except ImportError:
    GROQ_AVAILABLE = False



class AIService:
    def __init__(self):
        self.gemini_key = os.getenv("GEMINI_API_KEY")
        self.groq_key = os.getenv("GROQ_API_KEY")
        self.openai_key = os.getenv("OPENAI_API_KEY")
        self.api_key = self.gemini_key or self.groq_key or self.openai_key
        self.mode = "LLM_ENABLED" if self.api_key else "INTELLIGENT_SYNTHESIS"

    # =========================================================================
    # 1. SHORT ANSWER EVALUATION (RUBRIC & CONCEPT MATCHING)
    # =========================================================================
    def evaluate_short_answer(self, question_text: str, expected_answer: str, student_answer: str, max_marks: int = 2) -> dict:
        if not student_answer or not student_answer.strip():
            return {
                "marksObtained": 0.0,
                "maxMarks": max_marks,
                "percentage": 0.0,
                "feedback": "No answer provided.",
                "keywordsMatched": [],
                "missingConcepts": ["Complete response missing"]
            }

        student_norm = student_answer.lower().strip()
        expected_norm = expected_answer.lower().strip()

        stop_words = {"the", "a", "an", "is", "are", "was", "were", "in", "on", "of", "to", "and", "or", "that", "this", "it", "with", "by", "as", "for"}
        expected_words = [w for w in re.findall(r'\b[a-z0-9_]+\b', expected_norm) if w not in stop_words and len(w) > 2]
        student_words = set(re.findall(r'\b[a-z0-9_]+\b', student_norm))

        matched = [w for w in set(expected_words) if w in student_words]
        keyword_ratio = len(matched) / max(len(set(expected_words)), 1)
        seq_ratio = difflib.SequenceMatcher(None, expected_norm, student_norm).ratio()
        semantic_score = (keyword_ratio * 0.65) + (seq_ratio * 0.35)

        if semantic_score >= 0.75:
            awarded = float(max_marks)
            feedback = "Excellent conceptual explanation! Key principles and terminologies are clearly articulated."
        elif semantic_score >= 0.50:
            awarded = round(max_marks * 0.75, 1)
            feedback = "Good answer covering core concepts, but missing slight precision or key technical details."
        elif semantic_score >= 0.25:
            awarded = round(max_marks * 0.40, 1)
            feedback = "Partially correct. Identified some related ideas, but lacks depth and essential terminology."
        else:
            awarded = 0.0
            feedback = "Incorrect or insufficient conceptual grasp. Review recommended prerequisites."

        missing = [w for w in set(expected_words) if w not in student_words][:3]

        return {
            "marksObtained": awarded,
            "maxMarks": max_marks,
            "percentage": round((awarded / max_marks) * 100, 1),
            "feedback": feedback,
            "keywordsMatched": matched[:5],
            "missingConcepts": missing
        }

    # =========================================================================
    # 2. DIAGNOSTIC ASSESSMENT & LEARNING PATH SYNTHESIS
    # =========================================================================
    def analyze_diagnostic(self, responses: list) -> dict:
        topic_stats = {}
        for item in responses:
            topic = item.get("topic", "General DSA")
            is_correct = bool(item.get("isCorrect", False))
            if topic not in topic_stats:
                topic_stats[topic] = {"total": 0, "correct": 0}
            topic_stats[topic]["total"] += 1
            if is_correct:
                topic_stats[topic]["correct"] += 1

        topic_breakdown = []
        primary_gap = None
        min_acc = 101

        for topic, stat in topic_stats.items():
            acc = int(round((stat["correct"] / max(stat["total"], 1)) * 100))
            if acc >= 80:
                skill = "Strong"
                gap = "Low"
            elif acc >= 55:
                skill = "Medium"
                gap = "Medium"
            elif acc >= 35:
                skill = "Weak"
                gap = "High"
            else:
                skill = "Critical"
                gap = "Critical"

            topic_breakdown.append({
                "topic": topic,
                "accuracy": acc,
                "skillLevel": skill,
                "gap": gap,
                "correct": stat["correct"],
                "total": stat["total"]
            })

            if acc < min_acc:
                min_acc = acc
                primary_gap = {
                    "topic": topic,
                    "accuracy": acc,
                    "skillLevel": skill,
                    "gap": gap
                }

        if not topic_breakdown:
            topic_breakdown = [
                {"topic": "Arrays", "accuracy": 90, "skillLevel": "Strong", "gap": "Low"},
                {"topic": "Linked Lists", "accuracy": 71, "skillLevel": "Medium", "gap": "Medium"},
                {"topic": "Stacks", "accuracy": 88, "skillLevel": "Strong", "gap": "Low"},
                {"topic": "Queues", "accuracy": 68, "skillLevel": "Medium", "gap": "Medium"},
                {"topic": "Trees", "accuracy": 32, "skillLevel": "Weak", "gap": "High"},
                {"topic": "Graphs", "accuracy": 20, "skillLevel": "Critical", "gap": "Critical"},
                {"topic": "Sorting", "accuracy": 85, "skillLevel": "Strong", "gap": "Low"}
            ]
            primary_gap = {"topic": "Trees", "accuracy": 32, "skillLevel": "Weak", "gap": "High"}

        return {
            "topicBreakdown": topic_breakdown,
            "primaryGap": primary_gap or topic_breakdown[4],
            "recommendedFocus": f"Resolving {primary_gap['topic']} ({primary_gap['accuracy']}%) learning gap via a targeted 45-minute daily path"
        }

    # =========================================================================
    # 3. YOUTUBE / NPTEL / FILE CONTENT EXTRACTION & METADATA
    # =========================================================================
    def extract_youtube_id(self, url: str) -> Optional[str]:
        """Extracts 11-char YouTube video ID from various URL structures."""
        if not url:
            return None
        patterns = [
            r'(?:v=|\/)([0-9A-Za-z_-]{11}).*',
            r'(?:youtu\.be\/)([0-9A-Za-z_-]{11})',
            r'(?:embed\/)([0-9A-Za-z_-]{11})',
            r'(?:shorts\/)([0-9A-Za-z_-]{11})'
        ]
        for pattern in patterns:
            match = re.search(pattern, url)
            if match:
                return match.group(1)
        return None

    def fetch_youtube_metadata(self, video_id: str) -> dict:
        """Queries YouTube oEmbed endpoint to retrieve genuine title & channel."""
        try:
            oembed_url = f"https://www.youtube.com/oembed?url=https://www.youtube.com/watch?v={video_id}&format=json"
            req = urllib.request.Request(oembed_url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
            with urllib.request.urlopen(req, timeout=4) as response:
                data = json.loads(response.read().decode())
                return {
                    "title": data.get("title", "Video Lecture"),
                    "author": data.get("author_name", "Academic Educator"),
                    "thumbnail": data.get("thumbnail_url", f"https://img.youtube.com/vi/{video_id}/hqdefault.jpg"),
                    "valid": True
                }
        except Exception:
            return {
                "title": f"YouTube Lecture ({video_id})",
                "author": "Computer Science Educator",
                "thumbnail": f"https://img.youtube.com/vi/{video_id}/hqdefault.jpg",
                "valid": True
            }

    def fetch_youtube_transcript(self, video_id: str) -> tuple[Optional[str], bool]:
        """Attempts fetching English/auto captions using youtube_transcript_api."""
        if not YOUTUBE_TRANSCRIPT_AVAILABLE or not video_id:
            return None, False
        try:
            api = YouTubeTranscriptApi()
            transcript_list = api.fetch(video_id)
            full_text = " ".join([getattr(item, 'text', str(item)) for item in transcript_list])
            if len(full_text.strip()) > 40:
                return full_text, True
        except Exception as e1:
            try:
                if hasattr(YouTubeTranscriptApi, 'get_transcript'):
                    transcript_list = YouTubeTranscriptApi.get_transcript(video_id, languages=['en', 'en-US', 'en-IN', 'hi'])
                    full_text = " ".join([item.get('text', '') for item in transcript_list if isinstance(item, dict)])
                    if len(full_text.strip()) > 40:
                        return full_text, True
            except Exception as e2:
                print(f"Transcript fetch status for video {video_id}: {str(e1)} / {str(e2)}")
        return None, False


    def fetch_nptel_info(self, url: str) -> dict:
        """Parses NPTEL lecture page metadata and public transcript/text."""
        info = {"title": "NPTEL Academic Lecture", "channel": "NPTEL / IIT Ministry of Education", "content": ""}
        if not url:
            return info
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
            with urllib.request.urlopen(req, timeout=5) as resp:
                html = resp.read().decode('utf-8', errors='ignore')
                title_match = re.search(r'<title>(.*?)</title>', html, re.IGNORECASE)
                if title_match:
                    clean_title = re.sub(r'\s+', ' ', title_match.group(1)).strip()
                    info["title"] = clean_title
                desc_match = re.search(r'<meta\s+name=["\']description["\']\s+content=["\'](.*?)["\']', html, re.IGNORECASE)
                if desc_match:
                    info["content"] = desc_match.group(1)
                else:
                    paras = re.findall(r'<p[^>]*>(.*?)</p>', html, re.DOTALL | re.IGNORECASE)
                    clean_paras = [re.sub(r'<[^>]+>', '', p).strip() for p in paras if len(p.strip()) > 30]
                    info["content"] = "\n".join(clean_paras[:5])
        except Exception as e:
            print(f"NPTEL fetch status for {url}: {str(e)}")
        return info

    def extract_pdf_content(self, content_bytes: bytes) -> tuple[str, bool, str]:
        """Extracts text from PDF bytes using pypdf. Returns (text, is_valid, note)."""
        if not PYPDF_AVAILABLE:
            return "", False, "PDF processing library (pypdf) is not available."
        try:
            import io
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                txt = page.extract_text()
                if txt and txt.strip():
                    pages_text.append(f"--- Page / Section {idx + 1} ---\n" + txt.strip())
            full_text = "\n\n".join(pages_text)
            if len(full_text.strip()) < 50:
                return "", False, "The uploaded PDF contains scanned pages or image-only content with insufficient extractable text."
            return full_text, True, ""
        except Exception as e:
            return "", False, f"Could not read PDF document: {str(e)}"

    def extract_ppt_content(self, content_bytes: bytes) -> tuple[str, list, bool, str]:
        """Parses PPT/PPTX slide text, titles, bullet points, speaker notes using python-pptx."""
        if not PPTX_AVAILABLE:
            return "", [], False, "PowerPoint processing library (python-pptx) is not available."
        try:
            import io
            prs = Presentation(io.BytesIO(content_bytes))
            slides_data = []
            full_text_parts = []

            for idx, slide in enumerate(prs.slides):
                slide_num = idx + 1
                title = f"Slide {slide_num}"
                slide_texts = []

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                slide_texts.append(t)
                                if shape == slide.shapes[0] and len(t) < 100:
                                    title = t

                notes_text = ""
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()

                slides_data.append({
                    "slideNumber": slide_num,
                    "title": title,
                    "bulletPoints": slide_texts,
                    "notes": notes_text
                })

                slide_str = f"Slide {slide_num}: {title}\n" + "\n".join(slide_texts)
                if notes_text:
                    slide_str += f"\nSpeaker Notes: {notes_text}"
                full_text_parts.append(slide_str)

            full_text = "\n\n".join(full_text_parts)
            if len(full_text.strip()) < 20:
                return "", [], False, "Presentation appears to contain image slides with no extractable text."
            return full_text, slides_data, True, ""
        except Exception as e:
            return "", [], False, f"Could not read PowerPoint file: {str(e)}"


    # =========================================================================
    # 4. MODE-SPECIFIC PROMPT BUILDER & MULTI-PROVIDER LLM CALLER
    # =========================================================================
    def build_mode_prompt(self, input_mode: str, topic_label: str, content_text: str, extra_context: str = "") -> str:
        content_preview = content_text[:3500] if content_text else "Topic: " + topic_label

        if input_mode == "youtube":
            return f"""You are an AI Learning Resource Simplification Agent analyzing a YouTube video lecture for: "{topic_label}".

PRIMARY SOURCE / TRANSCRIPT PROVIDED BELOW:
{content_preview}

{extra_context}

CRITICAL RULES:
1. Ground your explanation PRIMARY on the actual transcript/content provided above.
2. DO NOT invent or hallucinate video content not supported by the transcript/metadata.
3. If transcript was unavailable, explicitly acknowledge this in the summary.
4. Preserve technical correctness while using student-friendly language.

Return ONLY a single valid JSON object with NO markdown formatting, matching this exact schema:
{{
  "topic": "{topic_label}",
  "inputMode": "youtube",
  "summary": "What this video/topic is about (2-3 sentences)",
  "simpleExplanation": "Complete easy explanation of the video content for students",
  "keyConcepts": ["Important Concept 1", "Important Concept 2", "Important Concept 3", "Important Concept 4"],
  "importantDefinitions": [
    {{"term": "Definition Term 1", "definition": "Clear plain language definition"}},
    {{"term": "Definition Term 2", "definition": "Clear plain language definition"}}
  ],
  "stepByStep": ["Step 1 walkthrough", "Step 2 walkthrough", "Step 3 walkthrough"],
  "examples": [
    {{"title": "Example from Video", "code": "Example code, algorithm, or text", "source": "source"}}
  ],
  "formulas": [
    {{"formula": "Important Formula / Equation", "explanation": "Explanation of formula"}}
  ],
  "examNotes": ["High-probability exam point 1", "High-probability exam point 2", "High-probability exam point 3"],
  "quickRevision": ["60-second revision point 1", "60-second revision point 2", "60-second revision point 3"],
  "practiceQuestions": [
    {{"question": "Practice Q1?", "answer": "Detailed answer Q1"}},
    {{"question": "Practice Q2?", "answer": "Detailed answer Q2"}},
    {{"question": "Practice Q3?", "answer": "Detailed answer Q3"}},
    {{"question": "Practice Q4?", "answer": "Detailed answer Q4"}},
    {{"question": "Practice Q5?", "answer": "Detailed answer Q5"}}
  ]
}}"""

        elif input_mode == "nptel":
            return f"""You are an AI Learning Resource Simplification Agent simplifying an NPTEL Academic Course Lecture for: "{topic_label}".

PRIMARY SOURCE / LECTURE CONTENT PROVIDED BELOW:
{content_preview}

{extra_context}

CRITICAL RULES:
1. Explain the NPTEL lecture in clear, student-friendly language suitable for college exam preparation.
2. Ground your notes in the available lecture content. Clearly indicate any limitations in source text.
3. Include technical formulas, step-by-step procedures, and high-level academic concepts.

Return ONLY a single valid JSON object matching this structure:
{{
  "topic": "{topic_label}",
  "inputMode": "nptel",
  "summary": "Course / Lecture Overview",
  "simpleExplanation": "Complete easy explanation of the available NPTEL lecture content.",
  "keyConcepts": ["Key Concept 1", "Key Concept 2", "Key Concept 3", "Key Concept 4"],
  "importantDefinitions": [
    {{"term": "Technical Term 1", "definition": "Plain language definition"}},
    {{"term": "Technical Term 2", "definition": "Plain language definition"}}
  ],
  "stepByStep": ["Step 1", "Step 2", "Step 3"],
  "formulas": [
    {{"formula": "Technical Formula / Equation", "explanation": "Explanation of formula"}}
  ],
  "examples": [
    {{"title": "Lecture Example", "code": "Example content", "source": "source"}}
  ],
  "examNotes": ["Exam important point 1", "Exam important point 2", "Exam important point 3"],
  "quickRevision": ["Quick revision bullet 1", "Quick revision bullet 2", "Quick revision bullet 3"],
  "practiceQuestions": [
    {{"question": "Exam Practice Q1?", "answer": "Detailed answer Q1"}},
    {{"question": "Exam Practice Q2?", "answer": "Detailed answer Q2"}},
    {{"question": "Exam Practice Q3?", "answer": "Detailed answer Q3"}},
    {{"question": "Exam Practice Q4?", "answer": "Detailed answer Q4"}},
    {{"question": "Exam Practice Q5?", "answer": "Detailed answer Q5"}}
  ]
}}"""

        elif input_mode in ["pdf", "file"]:
            return f"""You are an AI Learning Resource Simplification Agent simplifying an uploaded PDF / Textbook Chapter: "{topic_label}".

PRIMARY SOURCE (UPLOADED PDF TEXT):
{content_preview}

{extra_context}

CRITICAL RULES:
1. Simplify the SAME CONTENT present in the PDF.
2. DO NOT add unrelated textbook knowledge or change the original material's meaning.
3. DO NOT invent fake examples and present them as coming from the PDF.
4. If you provide additional AI-created examples, mark their source field as "ai_created". Original PDF examples must be marked "source".
5. Preserve headings, definitions, formulas, and steps found in the source text.

Return ONLY a single valid JSON object matching this structure:
{{
  "topic": "{topic_label}",
  "inputMode": "pdf",
  "summary": "1. Topic Overview: What the uploaded material is about.",
  "simpleExplanation": "2. Complete Simplified Explanation of the actual content from the uploaded document.",
  "keyConcepts": ["Key Concept 1 found in material", "Key Concept 2 found in material", "Key Concept 3 found in material"],
  "importantDefinitions": [
    {{"term": "Definition from material", "definition": "Plain language explanation"}}
  ],
  "formulas": [
    {{"formula": "Formula appearing in source", "explanation": "Explanation of formula"}}
  ],
  "examples": [
    {{"title": "Example from Uploaded Material", "code": "Example text", "source": "source"}},
    {{"title": "Extra AI Example (Supporting)", "code": "Extra supporting example", "source": "ai_created"}}
  ],
  "stepByStep": ["Step 1 from uploaded material", "Step 2 from uploaded material"],
  "examNotes": ["Important point from material 1", "Important point from material 2"],
  "quickRevision": ["Quick revision point 1", "Quick revision point 2"],
  "practiceQuestions": [
    {{"question": "Practice Question based on uploaded material?", "answer": "Answer based on material"}},
    {{"question": "Practice Question 2?", "answer": "Answer 2"}},
    {{"question": "Practice Question 3?", "answer": "Answer 3"}},
    {{"question": "Practice Question 4?", "answer": "Answer 4"}},
    {{"question": "Practice Question 5?", "answer": "Answer 5"}}
  ]
}}"""

        elif input_mode == "ppt":
            return f"""You are an AI Learning Resource Simplification Agent simplifying an uploaded PowerPoint Presentation (PPT) for: "{topic_label}".

PRIMARY SOURCE (SLIDE CONTENTS):
{content_preview}

{extra_context}

CRITICAL RULES:
1. Read the slides and provide a slide-by-slide simplified breakdown so the student understands the whole deck without reading original slides repeatedly.
2. Ground all explanations in the actual slide content.
3. Summarize overall concepts, formulas, definitions, and practice questions at the end.

Return ONLY a single valid JSON object matching this structure:
{{
  "topic": "{topic_label}",
  "inputMode": "ppt",
  "summary": "Complete Topic Overview of the Presentation",
  "simpleExplanation": "Unified simple explanation of all presentation slides.",
  "slideBySlide": [
    {{"slideNumber": 1, "title": "Slide 1 Topic", "simpleExplanation": "Simple explanation of Slide 1 content"}},
    {{"slideNumber": 2, "title": "Slide 2 Topic", "simpleExplanation": "Simple explanation of Slide 2 content"}}
  ],
  "keyConcepts": ["Key Concept 1 from slides", "Key Concept 2 from slides"],
  "importantDefinitions": [
    {{"term": "Term from slides", "definition": "Simple explanation"}}
  ],
  "formulas": [
    {{"formula": "Formula in slides", "explanation": "Explanation"}}
  ],
  "examNotes": ["High probability exam point 1", "High probability exam point 2"],
  "quickRevision": ["Quick revision bullet 1", "Quick revision bullet 2"],
  "practiceQuestions": [
    {{"question": "Question on slide material?", "answer": "Answer based on slides"}},
    {{"question": "Question 2?", "answer": "Answer 2"}},
    {{"question": "Question 3?", "answer": "Answer 3"}},
    {{"question": "Question 4?", "answer": "Answer 4"}},
    {{"question": "Question 5?", "answer": "Answer 5"}}
  ]
}}"""

        elif input_mode == "text":
            return f"""You are an AI Learning Resource Simplification Agent simplifying user-provided text for: "{topic_label}".

PRIMARY SOURCE (PASTED TEXT):
{content_preview}

CRITICAL RULES:
1. Use ONLY the supplied text as the primary source.
2. Preserve meaning, facts, definitions, formulas, and steps while making language easier.
3. Do not generate a generic explanation unrelated to the pasted text.

Return ONLY a single valid JSON object matching this structure:
{{
  "topic": "{topic_label}",
  "inputMode": "text",
  "summary": "1. What this text means (Overview)",
  "simpleExplanation": "2. Complete simplified explanation of the pasted text",
  "keyConcepts": ["Key Point 1 from text", "Key Point 2 from text"],
  "importantDefinitions": [
    {{"term": "Term in text", "definition": "Simple definition"}}
  ],
  "formulas": [
    {{"formula": "Formula in text", "explanation": "Explanation"}}
  ],
  "examples": [
    {{"title": "Example from text", "code": "Details", "source": "source"}}
  ],
  "examNotes": ["Exam Point 1 from text", "Exam Point 2 from text"],
  "quickRevision": ["Quick Revision 1", "Quick Revision 2"],
  "practiceQuestions": [
    {{"question": "Question based on pasted text?", "answer": "Answer based on text"}},
    {{"question": "Question 2?", "answer": "Answer 2"}},
    {{"question": "Question 3?", "answer": "Answer 3"}},
    {{"question": "Question 4?", "answer": "Answer 4"}},
    {{"question": "Question 5?", "answer": "Answer 5"}}
  ]
}}"""

        else:  # syllabus / topic
            return f"""You are an AI Learning Resource Simplification Agent teaching a computer science syllabus topic from scratch for university students: "{topic_label}".

{extra_context}

REQUIREMENTS:
1. Teach the topic completely, clearly, and comprehensively from fundamentals to advanced concepts.
2. Include a relatable real-life analogy, step-by-step breakdown, formulas/complexities, common student mistakes, text/Mermaid diagram descriptions, and practice Q&A.

Return ONLY a single valid JSON object matching this structure:
{{
  "topic": "{topic_label}",
  "inputMode": "syllabus",
  "summary": "1. Topic Overview & Plain Language Definition",
  "simpleExplanation": "2. Complete Concept Explanation with a relatable real-life analogy",
  "keyConcepts": ["Important Subtopic 1", "Important Subtopic 2", "Important Subtopic 3", "Important Subtopic 4"],
  "importantDefinitions": [
    {{"term": "Core Term 1", "definition": "Plain language definition"}},
    {{"term": "Core Term 2", "definition": "Plain language definition"}}
  ],
  "stepByStep": ["Step 1: How it works", "Step 2: Execution", "Step 3: Output/Completion"],
  "diagrams": [
    {{"title": "Visual Diagram / Flow", "mermaid": "graph TD\\nA[Start] --> B[Process] --> C[End]", "explanation": "Explanation of diagram"}}
  ],
  "examples": [
    {{"title": "Code / Algorithmic Implementation", "code": "Example code or step walkthrough", "source": "ai_created"}}
  ],
  "formulas": [
    {{"formula": "Time & Space Complexity / Math Formula", "explanation": "Explanation of complexity/formula"}}
  ],
  "examNotes": ["High probability exam question", "Key formula to remember"],
  "commonMistakes": ["Common Student Mistake 1", "Common Student Mistake 2"],
  "quickRevision": ["60-second summary point 1", "60-second summary point 2"],
  "practiceQuestions": [
    {{"question": "Practice Q1?", "answer": "Detailed solution for Q1"}},
    {{"question": "Practice Q2?", "answer": "Detailed solution for Q2"}},
    {{"question": "Practice Q3?", "answer": "Detailed solution for Q3"}},
    {{"question": "Practice Q4?", "answer": "Detailed solution for Q4"}},
    {{"question": "Practice Q5?", "answer": "Detailed solution for Q5"}}
  ]
}}"""

    def _call_llm_for_simplification(self, input_mode: str, resource_meta: dict, content: str, title: str, topic_hint: str, extra_context: str = "") -> Optional[dict]:
        topic_label = title or topic_hint or content[:100] or "Academic Learning Resource"
        prompt = self.build_mode_prompt(input_mode, topic_label, content, extra_context)

        # 1. Try Gemini API
        gemini_key = self.gemini_key or os.getenv("GEMINI_API_KEY")
        if GEMINI_AVAILABLE and gemini_key:
            try:
                genai.configure(api_key=gemini_key)
                for model_name in ['gemini-3.6-flash', 'gemini-3.5-flash', 'gemini-2.5-flash', 'gemini-2.0-flash', 'gemini-1.5-flash']:
                    try:
                        m = genai.GenerativeModel(model_name)
                        resp = m.generate_content(prompt)
                        if resp and resp.text:
                            raw_text = resp.text.strip()
                            clean_json = re.sub(r'^```json\s*', '', raw_text)
                            clean_json = re.sub(r'^```\s*', '', clean_json)
                            clean_json = re.sub(r'```$', '', clean_json.strip()).strip()
                            json_str = clean_json[clean_json.find('{'):clean_json.rfind('}')+1]
                            data = json.loads(json_str)
                            data["resourceMeta"] = resource_meta
                            return data
                    except Exception as ge:
                        print(f"Gemini model {model_name} note: {ge}")
            except Exception as e:
                print(f"Gemini API note: {e}")

        # 2. Try Groq REST / SDK API
        groq_key = self.groq_key or os.getenv("GROQ_API_KEY")
        if groq_key:
            models = ['groq/compound-mini', 'llama-3.3-70b-versatile', 'groq/compound']
            if GROQ_AVAILABLE:
                try:
                    client = Groq(api_key=groq_key)
                    for m in models:
                        try:
                            comp = client.chat.completions.create(
                                model=m,
                                messages=[
                                    {"role": "system", "content": "You output ONLY valid JSON matching the requested format."},
                                    {"role": "user", "content": prompt}
                                ],
                                temperature=0.2,
                                max_tokens=2000
                            )
                            if comp and comp.choices and comp.choices[0].message.content:
                                raw_text = comp.choices[0].message.content.strip()
                                clean_json = re.sub(r'^```json\s*', '', raw_text)
                                clean_json = re.sub(r'^```\s*', '', clean_json)
                                clean_json = re.sub(r'```$', '', clean_json.strip()).strip()
                                data = json.loads(clean_json[clean_json.find('{'):clean_json.rfind('}')+1])
                                data["resourceMeta"] = resource_meta
                                return data
                        except Exception:
                            continue
                except Exception:
                    pass

        return None

    def _build_mode_fallback(self, input_mode: str, resource_meta: dict, clean_content: str, title: str, topic_hint: str) -> dict:
        """Intelligent, topic-aware fallback when AI service is offline."""
        subject_text = f"{title} {topic_hint} {clean_content}".lower()
        if any(k in subject_text for k in ["tree", "binary tree", "bst", "traversal", "inorder", "preorder", "postorder", "avl"]):
            data = self._build_tree_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["graph", "bfs", "dfs", "dijkstra", "adjacency", "spanning"]):
            data = self._build_graph_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["stack", "lifo", "postfix", "infix", "parenthesis"]):
            data = self._build_stack_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["queue", "fifo", "circular queue", "priority queue", "deque"]):
            data = self._build_queue_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["linked list", "singly", "doubly", "pointer", "node"]):
            data = self._build_linked_list_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["sort", "quicksort", "mergesort", "binary search", "bubble"]):
            data = self._build_sorting_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["recursion", "recursive", "base case", "factorial"]):
            data = self._build_recursion_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["operating system", "os", "process", "thread", "deadlock", "paging"]):
            data = self._build_os_simplification(resource_meta, topic_hint)
        else:
            data = self._build_generic_simplification(resource_meta, clean_content, topic_hint)

        data["inputMode"] = input_mode
        data["resourceMeta"] = resource_meta
        return data

    def simplify_resource(self, content_type: str, content: str, title: str = "", topic_hint: str = "", metadata: Optional[dict] = None) -> dict:
        """
        AI Learning Resource Simplification Agent:
        Identifies input type (youtube, nptel, pdf, ppt, text, syllabus)
        and processes content according to input-type specific rules.
        """
        clean_content = (content or "").strip()
        raw_mode = (content_type or "syllabus").lower().strip()

        # Normalize inputMode
        if raw_mode in ["topic", "syllabus", "concept"]:
            input_mode = "syllabus"
        elif raw_mode in ["file", "pdf", "doc", "docx"]:
            if title.lower().endswith((".ppt", ".pptx")):
                input_mode = "ppt"
            else:
                input_mode = "pdf"
        elif raw_mode in ["ppt", "pptx"]:
            input_mode = "ppt"
        elif raw_mode in ["text", "notes"]:
            input_mode = "text"
        elif raw_mode == "nptel" or "nptel.ac.in" in clean_content.lower() or "nptel" in clean_content.lower():
            input_mode = "nptel"
        elif raw_mode == "youtube" or "youtube.com" in clean_content.lower() or "youtu.be" in clean_content.lower():
            input_mode = "youtube"
        else:
            input_mode = "syllabus"

        resource_meta = {
            "type": input_mode,
            "inputMode": input_mode,
            "sourceTitle": title or topic_hint or "Learning Resource",
            "sourceStatus": "normal"
        }

        extra_context = ""
        actual_text_for_prompt = clean_content

        # 1. YouTube URL
        if input_mode == "youtube":
            video_id = self.extract_youtube_id(clean_content)
            if video_id:
                meta = self.fetch_youtube_metadata(video_id)
                resource_meta["videoId"] = video_id
                resource_meta["youtubeUrl"] = f"https://www.youtube.com/watch?v={video_id}"
                resource_meta["sourceTitle"] = meta["title"]
                resource_meta["channel"] = meta["author"]
                resource_meta["thumbnail"] = meta["thumbnail"]
                topic_hint = topic_hint or meta["title"]

                transcript_text, has_transcript = self.fetch_youtube_transcript(video_id)
                if has_transcript:
                    resource_meta["sourceStatus"] = "transcript_available"
                    actual_text_for_prompt = f"Video Title: {meta['title']}\nChannel: {meta['author']}\nTranscript:\n{transcript_text}"
                else:
                    resource_meta["sourceStatus"] = "transcript_unavailable"
                    resource_meta["notesNotice"] = "Transcript could not be obtained for this YouTube video. Note: Simplification is based on video metadata and topic structure."
                    actual_text_for_prompt = f"Video Title: {meta['title']}\nChannel: {meta['author']}\n(Transcript unavailable for this video)"
                    extra_context = "NOTICE: Video transcript was unavailable. Clearly state in the summary that notes are derived from video metadata."
            else:
                resource_meta["sourceStatus"] = "invalid_url"
                resource_meta["notesNotice"] = "Invalid YouTube URL provided."

        # 2. NPTEL URL
        elif input_mode == "nptel":
            nptel_data = self.fetch_nptel_info(clean_content)
            resource_meta["sourceTitle"] = nptel_data.get("title", "NPTEL Video Lecture")
            resource_meta["channel"] = nptel_data.get("channel", "NPTEL / IIT Ministry of Education")
            topic_hint = topic_hint or nptel_data.get("title", "NPTEL Lecture")
            if nptel_data.get("content"):
                resource_meta["sourceStatus"] = "nptel_content_extracted"
                actual_text_for_prompt = f"Lecture Title: {nptel_data['title']}\nExtracted Content:\n{nptel_data['content']}"
            else:
                resource_meta["sourceStatus"] = "nptel_metadata_only"
                resource_meta["notesNotice"] = "NPTEL lecture page content was partially reachable. Notes cover the specified lecture topic."
                actual_text_for_prompt = f"NPTEL Lecture Topic: {title or topic_hint or clean_content}"

        # 3. PDF / Textbook / Notes
        elif input_mode == "pdf":
            if len(clean_content) < 50:
                resource_meta["sourceStatus"] = "pdf_unreadable"
                resource_meta["notesNotice"] = "The uploaded PDF appears to contain scanned pages or image-only content with insufficient extractable text."
                extra_context = "NOTICE: Extracted PDF text was sparse or scanned. Warn student that document text was partially unreadable."
            else:
                resource_meta["sourceStatus"] = "pdf_extracted"

        # 4. PowerPoint / PPT / PPTX
        elif input_mode == "ppt":
            if metadata and metadata.get("slides"):
                slides_info = metadata["slides"]
                slides_str = []
                for s in slides_info[:20]:
                    s_title = s.get("title", f"Slide {s.get('slideNumber')}")
                    bullets = "\n".join(s.get("bulletPoints", []))
                    notes = s.get("notes", "")
                    slides_str.append(f"--- Slide {s.get('slideNumber')}: {s_title} ---\n{bullets}\nNotes: {notes}")
                actual_text_for_prompt = "\n\n".join(slides_str)
                resource_meta["sourceStatus"] = "ppt_slides_extracted"
                resource_meta["slideCount"] = len(slides_info)
            else:
                resource_meta["sourceStatus"] = "ppt_text_extracted"

        # 5. Pasted Text
        elif input_mode == "text":
            resource_meta["sourceStatus"] = "text_provided"

        # 6. Syllabus Topic
        elif input_mode == "syllabus":
            resource_meta["sourceStatus"] = "syllabus_topic"

        # Generate via LLM
        ai_notes = self._call_llm_for_simplification(input_mode, resource_meta, actual_text_for_prompt, title, topic_hint, extra_context)
        if ai_notes:
            return ai_notes

        # Fallback if offline/failed
        return self._build_mode_fallback(input_mode, resource_meta, clean_content, title, topic_hint)


        # 2. Fallback Knowledge Base if AI is unavailable
        subject_text = f"{title} {topic_hint} {clean_content}".lower()
        if any(k in subject_text for k in ["tree", "binary tree", "bst", "traversal", "inorder", "preorder", "postorder", "avl"]):
            return self._build_tree_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["graph", "bfs", "dfs", "dijkstra", "adjacency", "spanning"]):
            return self._build_graph_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["stack", "lifo", "postfix", "infix", "parenthesis"]):
            return self._build_stack_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["queue", "fifo", "circular queue", "priority queue", "deque"]):
            return self._build_queue_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["linked list", "singly", "doubly", "pointer", "node"]):
            return self._build_linked_list_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["sort", "quicksort", "mergesort", "binary search", "bubble"]):
            return self._build_sorting_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["recursion", "recursive", "base case", "factorial"]):
            return self._build_recursion_simplification(resource_meta, topic_hint)
        elif any(k in subject_text for k in ["operating system", "os", "process", "thread", "deadlock", "paging"]):
            return self._build_os_simplification(resource_meta, topic_hint)
        else:
            return self._build_generic_simplification(resource_meta, clean_content, topic_hint)

    def _build_tree_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Binary Trees & Tree Traversals",
            "summary": "Trees are non-linear, hierarchical data structures where nodes are organized in parent-child relationships with exactly one root node and subtrees.",
            "simpleExplanation": "Imagine a family tree or an organizational chart. Instead of items lined up one after another like beads on a necklace (arrays or linked lists), items branch out downwards. The very top item is called the 'Root'. From each node, you can branch left or right. In a Binary Tree, each parent can have AT MOST 2 children. This branching structure allows fast searching, dividing problems into halves.",
            "keyConcepts": [
                "Root Node: The topmost node with no parent.",
                "Binary Tree Constraint: At most 2 children per node (left child and right child).",
                "Leaf Nodes: Bottom nodes that have 0 children (both left and right are NULL).",
                "Tree Height / Depth: The maximum number of edges from the root down to the farthest leaf.",
                "Binary Search Tree (BST) Property: Left subtree values < Root value < Right subtree values."
            ],
            "importantDefinitions": [
                {"term": "Binary Tree", "definition": "A hierarchical data structure where every parent node has at most two children, termed the left child and right child."},
                {"term": "Inorder Traversal", "definition": "A depth-first recursive walk visiting: Left Subtree -> Root -> Right Subtree. In a BST, this always yields nodes in sorted ascending order."},
                {"term": "Preorder Traversal", "definition": "A traversal visiting: Root -> Left Subtree -> Right Subtree. Frequently used to serialize, clone, or create prefix expressions."},
                {"term": "Postorder Traversal", "definition": "A traversal visiting: Left Subtree -> Right Subtree -> Root. Essential for safe bottom-up memory deallocation and postfix evaluation."}
            ],
            "importantPoints": [
                "The maximum number of nodes on level 'i' of a binary tree is 2^i (assuming root is level 0).",
                "A binary tree with height 'h' has at most 2^(h+1) - 1 nodes.",
                "Inorder traversal of any valid Binary Search Tree produces strictly ascending sorted numbers.",
                "Recursive traversal requires O(h) extra stack memory where h is tree height. For a balanced tree, h = log N; for a skewed tree, h = N."
            ],
            "examples": [
                {
                    "title": "Inorder Traversal Example",
                    "code": "// Given tree with Root 2, Left 1, Right 3\nvoid inorder(Node* root) {\n    if (root == NULL) return;\n    inorder(root->left);       // 1. Visit Left (1)\n    printf(\"%d \", root->data); // 2. Print Root (2)\n    inorder(root->right);      // 3. Visit Right (3)\n}\n// Output: 1 2 3 (Sorted order!)"
                }
            ],
            "stepByStep": [
                "Step 1: Start execution at the Root node.",
                "Step 2: Check base condition: If the current pointer is NULL, return immediately (unwind stack).",
                "Step 3: Recursively call the traversal function on the Left Child.",
                "Step 4: Process / print the current node's data.",
                "Step 5: Recursively call the traversal function on the Right Child.",
                "Step 6: Return control to the parent activation frame."
            ],
            "examNotes": [
                "Expected Question: Differentiate between Inorder, Preorder, and Postorder with formulas.",
                "Important Property: A binary tree cannot be uniquely reconstructed using only Inorder or only Preorder. You MUST have Inorder + (Preorder OR Postorder).",
                "Time Complexity of all three traversals: O(N) since every node is visited exactly once.",
                "Auxiliary Space Complexity: O(h) on the function call stack."
            ],
            "quickRevision": [
                "• Binary Tree: Max 2 children per parent.",
                "• Inorder: Left -> Root -> Right (Produces sorted order in BST).",
                "• Preorder: Root -> Left -> Right (Cloning / copying trees).",
                "• Postorder: Left -> Right -> Root (Deleting nodes from bottom-up).",
                "• Balanced Height: O(log N); Degenerate/Skewed Height: O(N)."
            ],
            "practiceQuestions": [
                {
                    "question": "What is the Inorder traversal sequence for a BST containing nodes with keys [40, 20, 60, 10, 30]?",
                    "answer": "10, 20, 30, 40, 60 (Since Inorder traversal of any BST always outputs nodes in sorted ascending order)."
                },
                {
                    "question": "Why is Postorder traversal preferred when deleting all nodes in a dynamically allocated binary tree?",
                    "answer": "Because Postorder visits both children before visiting the parent. Deallocating children first ensures parent pointers remain valid and prevents memory leaks / dangling pointer errors."
                },
                {
                    "question": "If a binary tree has N nodes, what is the total number of NULL pointers (leaves + single children)?",
                    "answer": "Exactly N + 1 NULL pointers. Every node has 2 pointers (2N total). N-1 of them point to child nodes; therefore, 2N - (N - 1) = N + 1 are NULL."
                }
            ]
        }

    def _build_graph_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Graph Data Structures & BFS / DFS Traversals",
            "summary": "Graphs are non-linear collections of vertices (nodes) connected by edges that can model networks, road systems, social connections, and circuit dependencies.",
            "simpleExplanation": "Unlike trees where there is a top root and no loops, graphs can connect any node to any other node. Think of cities connected by highways or friends on social media. Because you can have cycles (loops where you return to a starting city), you MUST keep track of a 'Visited' list so you don't get stuck in an endless loop!",
            "keyConcepts": [
                "Vertices (V): The nodes or data entities in the graph.",
                "Edges (E): The connections between pairs of vertices (can be directed or undirected).",
                "Adjacency Matrix: A 2D array matrix[V][V] where matrix[i][j] = 1 if edge exists. Space: O(V^2).",
                "Adjacency List: An array of linked lists/vectors where each index stores its neighbors. Space: O(V + E) — optimal for sparse graphs.",
                "Breadth First Search (BFS): Level-by-level exploration using a Queue.",
                "Depth First Search (DFS): Deep exploration along each branch before backtracking, using a Stack or recursion."
            ],
            "importantDefinitions": [
                {"term": "Graph", "definition": "A mathematical structure G = (V, E) consisting of a set of vertices V and a set of edges E linking pairs of vertices."},
                {"term": "Breadth First Search (BFS)", "definition": "A traversal algorithm that visits all neighbor vertices at the present depth level before moving to the next level vertices. Implemented via FIFO Queue."},
                {"term": "Depth First Search (DFS)", "definition": "A traversal algorithm that dives as deep as possible along each branch before backtracking. Implemented via recursive call stack or LIFO Stack."},
                {"term": "Connected Component", "definition": "A maximal subgraph in which any two vertices are connected to each other by paths."}
            ],
            "importantPoints": [
                "BFS always finds the SHORTEST path in an unweighted graph.",
                "DFS is used for topological sorting, cycle detection, and strongly connected components.",
                "Time complexity for both BFS and DFS using Adjacency List is O(V + E).",
                "Always maintain a `bool visited[V]` array to prevent infinite cycles."
            ],
            "examples": [
                {
                    "title": "BFS Algorithm Pseudocode",
                    "code": "void BFS(int startVertex) {\n    queue<int> q;\n    visited[startVertex] = true;\n    q.push(startVertex);\n    while(!q.empty()) {\n        int curr = q.front(); q.pop();\n        cout << curr << \" \";\n        for(int neighbor : adj[curr]) {\n            if(!visited[neighbor]) {\n                visited[neighbor] = true;\n                q.push(neighbor);\n            }\n        }\n    }\n}"
                }
            ],
            "stepByStep": [
                "Step 1: Pick a starting vertex and mark it as visited.",
                "Step 2: Insert the starting vertex into a FIFO Queue.",
                "Step 3: Dequeue the front vertex and process/print it.",
                "Step 4: Loop through all immediate neighbors of this vertex.",
                "Step 5: For any neighbor not yet marked visited: mark it visited and enqueue it.",
                "Step 6: Repeat Steps 3-5 until the Queue becomes completely empty."
            ],
            "examNotes": [
                "Exam Trap: BFS cannot find shortest paths on weighted graphs with different edge weights (use Dijkstra's Algorithm instead).",
                "Space Complexity: Adjacency matrix is O(V^2); Adjacency list is O(V + E). For sparse graphs with few edges, Adjacency List is far more memory efficient.",
                "Cycle Detection: In an undirected graph, a cycle exists if a neighbor is already visited AND is NOT the parent of the current vertex."
            ],
            "quickRevision": [
                "• G = (V, E): Vertices + Edges.",
                "• BFS uses QUEUE -> Finds shortest path in unweighted graphs.",
                "• DFS uses STACK / RECURSION -> Used for cycle detection & topological sort.",
                "• Time: O(V + E) with Adjacency List; O(V^2) with Adjacency Matrix.",
                "• Never forget the `visited` array to prevent infinite loops!"
            ],
            "practiceQuestions": [
                {
                    "question": "Which data structure is essential for implementing Breadth First Search (BFS)?",
                    "answer": "A FIFO (First-In, First-Out) Queue data structure."
                },
                {
                    "question": "What is the time complexity of BFS and DFS when the graph is represented using an Adjacency Matrix?",
                    "answer": "O(V^2), because for every vertex we must scan all V entries in its row to find adjacent neighbors."
                },
                {
                    "question": "Can BFS be used to find the shortest path between two vertices in a weighted graph?",
                    "answer": "No. Standard BFS only works for unweighted graphs (or graphs where all edge weights are identical). For positive weighted graphs, Dijkstra's algorithm must be used."
                }
            ]
        }

    def _build_stack_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Stacks Data Structure & Applications",
            "summary": "A stack is a linear data structure following the LIFO (Last In, First Out) principle, where insertions and deletions happen only at the top.",
            "simpleExplanation": "Think of a stack of plates in a cafeteria. You place new plates on top, and when someone takes a plate, they take from the very top. The last plate placed on top is the first one removed! This simple rule makes stacks the backbone of function calls, undo buttons, and parenthesis matching.",
            "keyConcepts": [
                "LIFO Principle: Last-In, First-Out order.",
                "Push Operation: Adds an element to the top. O(1).",
                "Pop Operation: Removes the top element. O(1).",
                "Peek/Top Operation: Inspects the top element without removing it. O(1).",
                "Stack Overflow: Trying to push to a full stack.",
                "Stack Underflow: Trying to pop from an empty stack."
            ],
            "importantDefinitions": [
                {"term": "Stack", "definition": "A restricted linear list where elements are inserted and deleted from only one designated end called the top."},
                {"term": "LIFO", "definition": "Last-In, First-Out protocol governing stack operations."},
                {"term": "Infix Expression", "definition": "An expression format where the operator is placed between operands (e.g. A + B)."},
                {"term": "Postfix (RPN) Expression", "definition": "An expression format where the operator follows its operands (e.g. A B +), eliminating the need for parentheses."}
            ],
            "importantPoints": [
                "All fundamental stack operations (Push, Pop, Peek, IsEmpty) execute in strictly O(1) constant time.",
                "Function call execution in programming languages relies on the runtime Call Stack.",
                "Balanced parenthesis validation is solved elegantly using a stack in O(N) time.",
                "Infix to Postfix conversion uses an operator precedence stack."
            ],
            "examples": [
                {
                    "title": "Parenthesis Balancing with Stack",
                    "code": "bool isBalanced(string s) {\n    stack<char> st;\n    for(char c : s) {\n        if(c == '(' || c == '{' || c == '[') st.push(c);\n        else {\n            if(st.empty()) return false;\n            char top = st.top(); st.pop();\n            if((c == ')' && top != '(') ||\n               (c == '}' && top != '{') ||\n               (c == ']' && top != '[')) return false;\n        }\n    }\n    return st.empty();\n}"
                }
            ],
            "stepByStep": [
                "Step 1: Check if the stack is full before pushing (Stack Overflow check).",
                "Step 2: Increment the `top` index by 1.",
                "Step 3: Assign the incoming element at `arr[top]`.",
                "Step 4: To pop, check if `top == -1` (Stack Underflow check).",
                "Step 5: Store or return `arr[top]` and decrement `top` by 1."
            ],
            "examNotes": [
                "High Frequency Exam Topic: Convert Infix to Postfix using operator precedence table (^ > *, / > +, -).",
                "Evaluating Postfix Expression: Push operands to stack. When an operator is encountered, pop two operands, compute, and push result back.",
                "Tower of Hanoi problem is solved recursively with an implicit call stack in 2^N - 1 moves."
            ],
            "quickRevision": [
                "• LIFO: Last In, First Out.",
                "• Push, Pop, Peek: All are O(1) time.",
                "• Overflow: Full stack push; Underflow: Empty stack pop.",
                "• Applications: Undo/Redo, Call Stack, Parenthesis check, Expression evaluation."
            ],
            "practiceQuestions": [
                {
                    "question": "What is the postfix form of the infix expression: (A + B) * C?",
                    "answer": "A B + C *"
                },
                {
                    "question": "What happens if you attempt to perform a Pop operation on an empty stack?",
                    "answer": "Stack Underflow condition occurs."
                },
                {
                    "question": "Which data structure is utilized internally to manage recursive function calls?",
                    "answer": "The system runtime Call Stack."
                }
            ]
        }

    def _build_queue_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Queues & Circular Queue Mechanics",
            "summary": "A queue is a linear structure following the FIFO (First In, First Out) principle, where elements enter at the rear and depart from the front.",
            "simpleExplanation": "Imagine standing in line at a movie ticket counter. The person who arrives first gets served first! New arrivals join the back (Rear), while the person who finished leaves from the front (Front). A Circular Queue solves the wasted space problem of linear arrays by wrapping the rear pointer back to 0 using modulo arithmetic.",
            "keyConcepts": [
                "FIFO Principle: First-In, First-Out order.",
                "Enqueue: Insert an element at the Rear end. O(1).",
                "Dequeue: Remove an element from the Front end. O(1).",
                "Circular Queue: Modulo index wrapping `(rear + 1) % capacity`.",
                "Priority Queue: Elements dequeued based on priority rather than arrival order."
            ],
            "importantDefinitions": [
                {"term": "Queue", "definition": "A linear collection with two open ends: elements enter via the rear and leave via the front."},
                {"term": "FIFO", "definition": "First-In, First-Out principle governing queue access."},
                {"term": "Circular Queue", "definition": "A queue where the last position is connected back to the first position to make a circle, eliminating wasted memory space."},
                {"term": "Deque (Double-Ended Queue)", "definition": "A generalized queue allowing insertion and deletion at both front and rear ends."}
            ],
            "importantPoints": [
                "Linear array queues suffer from 'false overflow' when front moves forward and space behind it is unusable.",
                "Circular Queue uses `(index + 1) % size` to reuse freed front slots.",
                "Full condition in circular queue: `(rear + 1) % size == front`.",
                "Empty condition in circular queue: `front == -1`."
            ],
            "examples": [
                {
                    "title": "Circular Queue Enqueue Formula",
                    "code": "void enqueue(int val) {\n    if ((rear + 1) % SIZE == front) {\n        cout << \"Queue is Full! (Overflow)\";\n        return;\n    }\n    if (front == -1) front = 0; // First element\n    rear = (rear + 1) % SIZE;\n    arr[rear] = val;\n}"
                }
            ],
            "stepByStep": [
                "Step 1: Check if circular queue is full: `(rear + 1) % size == front`.",
                "Step 2: If empty (`front == -1`), set `front = 0`.",
                "Step 3: Update rear pointer: `rear = (rear + 1) % size`.",
                "Step 4: Store value at `arr[rear]`.",
                "Step 5: For dequeue, retrieve `arr[front]`.",
                "Step 6: If `front == rear`, queue has become empty: reset `front = rear = -1`."
            ],
            "examNotes": [
                "Exam Derivation: Show why Circular Queue solves linear queue false overflow.",
                "Number of elements in circular queue: `(rear - front + capacity) % capacity + 1`.",
                "Applications: CPU round-robin task scheduling, printer spooling, BFS graph traversal."
            ],
            "quickRevision": [
                "• FIFO: First In, First Out.",
                "• Enqueue at Rear, Dequeue at Front.",
                "• Circular Queue uses Modulo: `(rear + 1) % SIZE`.",
                "• Applications: BFS, CPU scheduling, buffer caching."
            ],
            "practiceQuestions": [
                {
                    "question": "What condition indicates that a Circular Queue of capacity N is completely full?",
                    "answer": "(rear + 1) % N == front"
                },
                {
                    "question": "Why is a circular queue preferred over a standard linear array queue?",
                    "answer": "Because a standard linear array queue cannot reuse freed slots after elements are dequeued from the front (false overflow), whereas a circular queue reclaims those spaces using modulo arithmetic."
                },
                {
                    "question": "Which traversal algorithm in graphs inherently depends on a Queue?",
                    "answer": "Breadth First Search (BFS)."
                }
            ]
        }

    def _build_linked_list_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Linked Lists (Singly, Doubly, Circular)",
            "summary": "A linked list is a linear data structure of nodes where each node stores data and a pointer/reference to the next node in dynamic memory.",
            "simpleExplanation": "Unlike arrays where items must sit side-by-side in continuous memory slots, linked list nodes can live anywhere in memory! Each node has two parts: the data it holds, and a pointer that points to where the next node lives. Inserting or deleting an element simply requires changing pointer arrows without shifting any other items!",
            "keyConcepts": [
                "Dynamic Size: Grows and shrinks at runtime without memory reallocation.",
                "Singly Linked List: Each node points only to the next node (`node->next`).",
                "Doubly Linked List: Each node points to both next and previous nodes (`node->prev`, `node->next`).",
                "Circular Linked List: The last node points back to the head node instead of NULL.",
                "Head Pointer: Stores the address of the first node.",
                "O(1) Insertion at Head: No shifting needed."
            ],
            "importantDefinitions": [
                {"term": "Node", "definition": "The basic building block of a linked list containing data and one or more pointer references."},
                {"term": "Head", "definition": "A reference pointer pointing to the first node in the linked list."},
                {"term": "Null Pointer", "definition": "A special pointer value in the last node indicating the termination of the list."},
                {"term": "Doubly Linked List", "definition": "A linked list where each node maintains two pointers: one to the forward neighbor and one to the backward neighbor."}
            ],
            "importantPoints": [
                "Arrays offer O(1) random access; Linked Lists require O(N) sequential traversal to access index i.",
                "Insertion/deletion at the beginning of a linked list is O(1) constant time.",
                "Reversing a singly linked list requires 3 pointers: `prev`, `curr`, and `next`.",
                "Always check for `head == NULL` to avoid segmentation faults (NullPointerExceptions)."
            ],
            "examples": [
                {
                    "title": "Reverse a Singly Linked List",
                    "code": "Node* reverseList(Node* head) {\n    Node* prev = NULL;\n    Node* curr = head;\n    while (curr != NULL) {\n        Node* nextTemp = curr->next; // 1. Save next\n        curr->next = prev;           // 2. Reverse pointer\n        prev = curr;                 // 3. Move prev forward\n        curr = nextTemp;             // 4. Move curr forward\n    }\n    return prev; // New head\n}"
                }
            ],
            "stepByStep": [
                "Step 1: To insert at head: allocate a new node with data.",
                "Step 2: Set `newNode->next = head`.",
                "Step 3: Update `head = newNode`.",
                "Step 4: To delete a node: traverse until finding the target node while tracking the previous node.",
                "Step 5: Set `prev->next = target->next`.",
                "Step 6: Free the target node's allocated memory."
            ],
            "examNotes": [
                "Classic Exam Problem: Detect a cycle/loop in a linked list using Floyd's Cycle-Finding Algorithm (Slow and Fast pointer).",
                "Comparison: Arrays have cache locality and fast access; Linked Lists have fast insertions/deletions without shifting.",
                "Memory Overhead: Each node requires extra bytes for pointer storage."
            ],
            "quickRevision": [
                "• Node = Data + Next pointer.",
                "• Insertion at Head: O(1); Access at Index: O(N).",
                "• Reverse: Use 3 pointers (prev, curr, next).",
                "• Cycle Detection: Floyd's Tortoise and Hare (slow by 1, fast by 2)."
            ],
            "practiceQuestions": [
                {
                    "question": "What is the time complexity to insert a new node at the beginning of a Singly Linked List?",
                    "answer": "O(1) constant time, because you only update the new node's next pointer to point to head and reset head."
                },
                {
                    "question": "How does Floyd's Cycle-Finding Algorithm detect a loop in a linked list?",
                    "answer": "By using two pointers: a slow pointer moving 1 step at a time, and a fast pointer moving 2 steps. If a cycle exists, the fast pointer will eventually catch up and equal the slow pointer."
                },
                {
                    "question": "Why can't Binary Search be applied efficiently directly on a Singly Linked List in O(log N) time?",
                    "answer": "Because linked lists do not support O(1) random index access. Finding the middle element requires O(N) sequential pointer traversal, making binary search degrade to O(N)."
                }
            ]
        }

    def _build_sorting_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Sorting & Searching Algorithms",
            "summary": "Sorting arranges items in ordered sequence (ascending/descending), while searching finds target values efficiently within collections.",
            "simpleExplanation": "Imagine looking up a word in a dictionary. If the dictionary was in random order, you would have to check every single page one by one (Linear Search O(N)). But because it is sorted alphabetically, you can open right in the middle, see if your word comes before or after, and eliminate half the book with each step (Binary Search O(log N))!",
            "keyConcepts": [
                "Binary Search: Divide-and-conquer search on sorted arrays. O(log N) time.",
                "Quicksort: Partition around a pivot element. Average O(N log N), Worst O(N^2).",
                "Mergesort: Divide into halves, recursively sort, and merge. Guaranteed O(N log N).",
                "Stability: A sorting algorithm is stable if it preserves relative order of duplicate keys.",
                "In-Place Sorting: Uses O(1) auxiliary space (e.g. Quicksort, Heapsort)."
            ],
            "importantDefinitions": [
                {"term": "Binary Search", "definition": "An algorithm that finds the position of a target value within a sorted array by repeatedly dividing the search interval in half."},
                {"term": "Quicksort", "definition": "A divide-and-conquer algorithm that selects a pivot element and partitions the array such that smaller elements precede it and larger follow it."},
                {"term": "Mergesort", "definition": "A stable divide-and-conquer algorithm that divides the array into halves, sorts them recursively, and merges the sorted halves."},
                {"term": "Algorithm Stability", "definition": "The property where elements with identical keys appear in the same relative order in the output as in the input."}
            ],
            "importantPoints": [
                "Binary Search strictly requires the input array to be pre-sorted.",
                "Comparison-based sorting has a mathematical lower bound of O(N log N).",
                "Mergesort requires O(N) extra temporary space for merging; Quicksort is in-place.",
                "Worst case for Quicksort occurs when the array is already sorted and the pivot chosen is always the first or last element."
            ],
            "examples": [
                {
                    "title": "Binary Search Iterative Implementation",
                    "code": "int binarySearch(int arr[], int n, int key) {\n    int low = 0, high = n - 1;\n    while (low <= high) {\n        int mid = low + (high - low) / 2;\n        if (arr[mid] == key) return mid;\n        else if (arr[mid] < key) low = mid + 1;\n        else high = mid - 1;\n    }\n    return -1; // Not found\n}"
                }
            ],
            "stepByStep": [
                "Step 1: Set `low = 0` and `high = n - 1`.",
                "Step 2: Calculate middle index: `mid = low + (high - low) / 2` (prevents integer overflow).",
                "Step 3: If `arr[mid] == target`, return `mid`.",
                "Step 4: If `arr[mid] < target`, discard left half: set `low = mid + 1`.",
                "Step 5: If `arr[mid] > target`, discard right half: set `high = mid - 1`.",
                "Step 6: Repeat until `low > high`. If not found, return -1."
            ],
            "examNotes": [
                "Exam Question: Compare Quicksort vs Mergesort (Time complexity, space complexity, stability, in-place).",
                "Mid Calculation: Explain why `mid = low + (high - low)/2` is safer than `(low + high)/2` in C/C++.",
                "Non-comparison sorts like Counting Sort and Radix Sort can achieve O(N) time under key range constraints."
            ],
            "quickRevision": [
                "• Binary Search: O(log N) — Must be SORTED.",
                "• Mergesort: Guaranteed O(N log N), Stable, O(N) extra space.",
                "• Quicksort: Average O(N log N), In-place, Worst O(N^2).",
                "• Best general-purpose: Hybrid (Timsort, Introsort)."
            ],
            "practiceQuestions": [
                {
                    "question": "What is the maximum number of comparisons required to binary search an element in a sorted array of 1,000,000 elements?",
                    "answer": "At most 20 comparisons, since 2^20 = 1,048,576 > 1,000,000 (ceil(log2(1,000,000)) = 20)."
                },
                {
                    "question": "Why is Mergesort preferred over Quicksort for sorting linked lists?",
                    "answer": "Because linked list nodes can be merged in-place in O(1) auxiliary space without random access, and Mergesort does not suffer from worst-case O(N^2) behavior."
                },
                {
                    "question": "What pivot selection strategy avoids Quicksort's worst-case O(N^2) time on already sorted arrays?",
                    "answer": "Randomized pivot selection or Median-of-Three pivot selection (taking the median of the first, middle, and last elements)."
                }
            ]
        }

    def _build_recursion_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Recursion & Call Stack Mechanics",
            "summary": "Recursion is a problem-solving technique where a function solves a problem by calling itself with a smaller subproblem until reaching a base termination case.",
            "simpleExplanation": "Think of Russian nesting dolls. You open the big doll to find a slightly smaller doll inside. You keep opening smaller dolls until you hit the tiny solid doll at the center that cannot open (the Base Case). Once you reach the center, you close the dolls back up one by one (Call Stack Unwinding). Every recursive function needs a base case, or it runs forever until your computer crashes with a Stack Overflow!",
            "keyConcepts": [
                "Base Case: The stopping condition that terminates recursion without further calls.",
                "Recursive Step: The function calling itself with a reduced, smaller input.",
                "Call Stack: Memory frames holding local variables and return addresses for each active call.",
                "Stack Unwinding: The return phase where values are calculated backwards.",
                "Tail Recursion: When the recursive call is the very last statement in the function."
            ],
            "importantDefinitions": [
                {"term": "Recursion", "definition": "A programming methodology where a function calls itself directly or indirectly to solve smaller instances of the same problem."},
                {"term": "Base Case", "definition": "The terminating condition that stops recursive execution and prevents infinite loops."},
                {"term": "Stack Overflow", "definition": "A runtime crash occurring when the function call stack exceeds its allotted memory due to missing or unreachable base cases."},
                {"term": "Recurrence Relation", "definition": "A mathematical equation defining a sequence based on its earlier terms (e.g. T(N) = 2T(N/2) + O(N))."}
            ],
            "importantPoints": [
                "Every recursive algorithm can be converted to an iterative one using an explicit stack.",
                "Without a base case, recursion leads directly to Stack Overflow.",
                "Each recursive call consumes extra stack memory proportional to the depth of recursion.",
                "Master Theorem provides an easy way to solve divide-and-conquer recurrence relations."
            ],
            "examples": [
                {
                    "title": "Factorial with Call Stack Trace",
                    "code": "int factorial(int n) {\n    if (n <= 1) return 1; // Base case!\n    return n * factorial(n - 1); // Recursive call\n}\n// Trace for factorial(3):\n// factorial(3) = 3 * factorial(2)\n// factorial(2) = 2 * factorial(1)\n// factorial(1) = 1 (Base case reached, unwinding!)\n// Returns: 2 * 1 = 2 -> 3 * 2 = 6"
                }
            ],
            "stepByStep": [
                "Step 1: Identify the smallest subproblem with a known, direct answer (Base Case).",
                "Step 2: Write the `if (base_condition) return base_value;` at the very top.",
                "Step 3: Break the main problem into one or more identical smaller subproblems.",
                "Step 4: Make the recursive call with the smaller parameter (e.g. `n - 1` or `n / 2`).",
                "Step 5: Combine the returned results from subproblems to compute the current answer.",
                "Step 6: Return the computed result to the caller."
            ],
            "examNotes": [
                "Exam Question: Solve recurrence relations using Master Theorem (T(N) = aT(N/b) + f(N)).",
                "Tail Call Optimization (TCO): Modern compilers can optimize tail-recursive functions to run in O(1) space like a while-loop.",
                "Contrast: Recursion is elegant and readable; Iteration is often faster with zero stack memory overhead."
            ],
            "quickRevision": [
                "• Recursion: Function calling itself with smaller input.",
                "• MUST have a BASE CASE to stop.",
                "• Memory: Uses function call stack (danger of Stack Overflow).",
                "• Unwinding: Evaluates answers on the way back up."
            ],
            "practiceQuestions": [
                {
                    "question": "What is the consequence of omitting a base case in a recursive function?",
                    "answer": "The function calls itself indefinitely until the memory reserved for the call stack is exhausted, causing a Stack Overflow runtime crash."
                },
                {
                    "question": "What is a Tail Recursive function?",
                    "answer": "A recursive function where the recursive call is the very last operation executed in the function with no pending computations."
                },
                {
                    "question": "Solve the recurrence relation T(N) = 2T(N/2) + O(N) using Master Theorem.",
                    "answer": "T(N) = O(N log N) (Standard Mergesort complexity, where a=2, b=2, log_b(a)=1, and f(N) = O(N^1))."
                }
            ]
        }

    def _build_os_simplification(self, meta: dict, hint: str) -> dict:
        return {
            "resourceMeta": meta,
            "topic": "Operating Systems & Process Management",
            "summary": "Operating Systems manage computer hardware resources, process execution, CPU scheduling, and virtual memory allocation.",
            "simpleExplanation": "Think of an Operating System as the air traffic controller of your computer. Multiple applications (processes) want to run on the CPU, write to memory, and read from disk simultaneously. The OS schedules which process gets the CPU, ensures one app cannot overwrite another app's memory, and gives every program the illusion of having the entire computer to itself.",
            "keyConcepts": [
                "Process vs Thread: A process is an independent executing program with its own memory space; a thread is a lightweight execution unit sharing process memory.",
                "CPU Scheduling: Algorithms (FCFS, SJF, Round Robin, Priority) that decide which ready process gets the CPU.",
                "Virtual Memory & Paging: Splitting memory into fixed-size pages and mapping virtual addresses to physical RAM frames.",
                "Deadlock: A state where two or more processes are blocked indefinitely, each holding a resource needed by another.",
                "4 Coffman Deadlock Conditions: Mutual Exclusion, Hold & Wait, No Preemption, Circular Wait."
            ],
            "importantDefinitions": [
                {"term": "Process", "definition": "A program in execution, including program counter, stack, and data section."},
                {"term": "Thread", "definition": "A lightweight unit of CPU execution within a process sharing the code, data, and open file descriptors."},
                {"term": "Deadlock", "definition": "A situation where a set of processes are blocked because each process is holding a resource and waiting for another resource held by another process."},
                {"term": "Paging", "definition": "A memory management scheme that eliminates the need for contiguous physical memory allocation by dividing virtual memory into pages."}
            ],
            "importantPoints": [
                "Context switching involves saving the CPU register state of the current process and loading the state of the next process (pure overhead).",
                "Round Robin scheduling uses a time quantum; if quantum is too small, excessive context switching occurs; if too large, it degrades to FCFS.",
                "Banker's Algorithm is used for Deadlock Avoidance by testing for safe states.",
                "Page Fault occurs when a program tries to access a page that is mapped in virtual memory but not currently loaded into physical RAM."
            ],
            "examples": [
                {
                    "title": "Round Robin CPU Scheduling Concept",
                    "code": "// Given processes P1(burst 6), P2(burst 4) with Quantum = 3\n// Time 0-3: P1 executes for 3 units (remaining: 3)\n// Time 3-6: P2 executes for 3 units (remaining: 1)\n// Time 6-9: P1 completes remaining 3 units\n// Time 9-10: P2 completes remaining 1 unit"
                }
            ],
            "stepByStep": [
                "Step 1: Process is created and placed in the 'Ready Queue'.",
                "Step 2: CPU Scheduler dispatches the process to the 'Running' state.",
                "Step 3: If an I/O request occurs, the process moves to 'Waiting/Blocked' state.",
                "Step 4: Once I/O finishes, it transitions back to the 'Ready' state.",
                "Step 5: When execution finishes, the process moves to the 'Terminated' state."
            ],
            "examNotes": [
                "High Frequency Exam Topic: 4 Necessary Conditions for Deadlock (Mutual Exclusion, Hold and Wait, No Preemption, Circular Wait).",
                "Process Synchronization: Peterson's Solution, Semaphores (Wait/Signal), and Mutex Locks.",
                "Belady's Anomaly: In FIFO page replacement, increasing the number of page frames can sometimes increase the number of page faults."
            ],
            "quickRevision": [
                "• Process = Program in execution (heavyweight, isolated memory).",
                "• Thread = Lightweight process (shares memory space).",
                "• Deadlock: 4 conditions (Mutual exclusion, Hold&wait, No preemption, Circular wait).",
                "• Scheduling: Round Robin for time-sharing, SJF for minimum average waiting time."
            ],
            "practiceQuestions": [
                {
                    "question": "What are the four Coffman conditions necessary for a deadlock to occur?",
                    "answer": "1. Mutual Exclusion, 2. Hold and Wait, 3. No Preemption, 4. Circular Wait."
                },
                {
                    "question": "What is the primary difference between a Process and a Thread?",
                    "answer": "A process has its own isolated memory address space and system resources, whereas threads exist within a process and share the same memory space, code section, and global variables."
                },
                {
                    "question": "What is Belady's Anomaly in operating systems?",
                    "answer": "The phenomenon in the FIFO page replacement algorithm where increasing the number of page frames leads to an unexpected increase in the number of page faults."
                }
            ]
        }

    def _build_generic_simplification(self, meta: dict, content: str, hint: str) -> dict:
        topic_name = hint or meta.get("sourceTitle") or "Academic Concept"
        preview = (content[:150] + "...") if len(content) > 150 else content
        return {
            "resourceMeta": meta,
            "topic": topic_name,
            "summary": f"Simplified study breakdown for {topic_name}. Formatted for quick comprehension and exam readiness.",
            "simpleExplanation": f"This material covers {topic_name}. In simple terms, this concept establishes fundamental rules and operations designed to organize information efficiently, optimize computer execution time, and eliminate redundant operations.",
            "keyConcepts": [
                f"Core Principle: Structured representation of {topic_name}.",
                "Efficiency: Minimizing execution steps and memory footprint.",
                "Modularity: Separating operations into clean, reusable procedures.",
                "Boundary Conditions: Handling empty states, edge inputs, and overflow conditions."
            ],
            "importantDefinitions": [
                {"term": topic_name, "definition": f"A foundational academic topic focusing on computational efficiency and structured logic."},
                {"term": "Time Complexity", "definition": "A mathematical representation of the number of operations executed as input size scales."},
                {"term": "Space Complexity", "definition": "The memory volume required by an algorithm to execute to completion."}
            ],
            "importantPoints": [
                "Always check edge cases (e.g. empty inputs, single element, negative values).",
                "Understand the trade-off between execution speed and memory consumption.",
                "Formulate a step-by-step trace before writing code.",
                "Review practice questions regularly to reinforce conceptual memory."
            ],
            "examples": [
                {
                    "title": f"{topic_name} Implementation Pattern",
                    "code": f"// Standard template for {topic_name}\nvoid process() {{\n    // 1. Initialize data\n    // 2. Execute logic\n    // 3. Return verified output\n}}"
                }
            ],
            "stepByStep": [
                "Step 1: Parse the input and check all prerequisite requirements.",
                "Step 2: Initialize required tracking variables and memory pointers.",
                "Step 3: Execute the core operational sequence.",
                "Step 4: Verify boundary conditions and finalize output."
            ],
            "examNotes": [
                f"Key Exam Focus: Define {topic_name}, draw the diagram/flowchart, and state the Time & Space complexities.",
                "Always write down the base case and boundary condition checks in written exams to score full marks."
            ],
            "quickRevision": [
                f"• Focus on core definition of {topic_name}.",
                "• Remember the primary operational time complexity.",
                "• Master the standard edge case handling steps."
            ],
            "practiceQuestions": [
                {
                    "question": f"What is the primary advantage of utilizing {topic_name}?",
                    "answer": "It provides a standardized, optimized approach to solving problems efficiently with minimal computational overhead."
                },
                {
                    "question": "What is the first step you should verify when processing any algorithm?",
                    "answer": "Check boundary/edge cases (such as null pointers, empty data structures, or index out of range)."
                }
            ]
        }

    # =========================================================================
    # 6. DEFAULT IMPORTANT NOTES FOR YOU (CURRICULUM BASED)
    # =========================================================================
    def get_default_important_notes(self, department: str = "CSE", year: str = "2") -> list:
        """
        Returns structured syllabus-aligned default study notes for Data Structures students
        (Arrays, Linked Lists, Stacks, Queues, Trees, Graphs, Sorting, Recursion).
        """
        return [
            {
                "id": "note-trees",
                "topic": "Binary Trees & Traversal Order",
                "department": "CSE",
                "badge": "Unit III • High Priority",
                "shortSummary": "Trees branch out hierarchically. Inorder (L-N-R) gives sorted keys in BST; Postorder (L-R-N) is used for bottom-up node deletion.",
                "readTime": "4 min read",
                "keyFormula": "Inorder = Left -> Root -> Right",
                "simplification": self._build_tree_simplification({"type": "syllabus", "sourceTitle": "Unit 3: Binary Trees"}, "Binary Trees")
            },
            {
                "id": "note-graphs",
                "topic": "Graph BFS vs DFS Exploration",
                "department": "CSE",
                "badge": "Unit IV • Core Exam Topic",
                "shortSummary": "BFS uses a Queue to find shortest paths in unweighted graphs; DFS uses a Stack/Recursion for cycle detection and topological sorting.",
                "readTime": "5 min read",
                "keyFormula": "BFS = Queue; DFS = Stack/Recursion",
                "simplification": self._build_graph_simplification({"type": "syllabus", "sourceTitle": "Unit 4: Graph Algorithms"}, "Graphs")
            },
            {
                "id": "note-stacks",
                "topic": "Stack LIFO & Expression Evaluation",
                "department": "CSE",
                "badge": "Unit II • Fundamental",
                "shortSummary": "Stacks enforce Last-In, First-Out (LIFO). Used for function call stack frames, balanced parenthesis checking, and infix-to-postfix conversion.",
                "readTime": "3 min read",
                "keyFormula": "Push/Pop/Peek = O(1) Constant Time",
                "simplification": self._build_stack_simplification({"type": "syllabus", "sourceTitle": "Unit 2: Stacks"}, "Stacks")
            },
            {
                "id": "note-queues",
                "topic": "Circular Queue Modulo Arithmetic",
                "department": "CSE",
                "badge": "Unit II • Exam Formula",
                "shortSummary": "Linear array queues cause false overflow. Circular queues reclaim freed front space using `(rear + 1) % capacity == front`.",
                "readTime": "4 min read",
                "keyFormula": "NextIndex = (index + 1) % Capacity",
                "simplification": self._build_queue_simplification({"type": "syllabus", "sourceTitle": "Unit 2: Queues"}, "Queues")
            },
            {
                "id": "note-linked-lists",
                "topic": "Linked Lists & Pointer Manipulation",
                "department": "CSE",
                "badge": "Unit I • Core",
                "shortSummary": "Linked lists provide dynamic memory allocation without contiguous space. In-place reversal uses 3 pointers (prev, curr, next).",
                "readTime": "4 min read",
                "keyFormula": "Reverse: curr->next = prev",
                "simplification": self._build_linked_list_simplification({"type": "syllabus", "sourceTitle": "Unit 1: Linked Lists"}, "Linked Lists")
            },
            {
                "id": "note-sorting",
                "topic": "Quicksort vs Mergesort & Binary Search",
                "department": "CSE",
                "badge": "Unit V • Algorithms",
                "shortSummary": "Binary search is O(log N) on sorted data. Quicksort is in-place average O(N log N); Mergesort is guaranteed O(N log N) and stable.",
                "readTime": "5 min read",
                "keyFormula": "Binary Search: O(log N)",
                "simplification": self._build_sorting_simplification({"type": "syllabus", "sourceTitle": "Unit 5: Sorting & Searching"}, "Sorting")
            },
            {
                "id": "note-recursion",
                "topic": "Recursion & Call Stack Unwinding",
                "department": "CSE",
                "badge": "Foundation • Prerequisite",
                "shortSummary": "Every recursive function MUST have a base case to terminate, or it crashes with Stack Overflow. Unwinding computes results backwards.",
                "readTime": "3 min read",
                "keyFormula": "Base Case + Recursive Step",
                "simplification": self._build_recursion_simplification({"type": "syllabus", "sourceTitle": "Foundations: Recursion"}, "Recursion")
            }
        ]

    # =========================================================================
    # 7. DYNAMIC CONTEXT-AWARE AI CHATBOT (NO PREDEFINED CANNED ANSWERS)
    # =========================================================================
    def _fetch_wiki_summary(self, term: str) -> Optional[Dict[str, str]]:
        try:
            clean = re.sub(r'^(what is|explain|tell me about|how does|why do we need|describe|meaning of)\s+', '', term.strip(), flags=re.IGNORECASE)
            clean = re.sub(r'[?!.]+$', '', clean).strip()
            if not clean or len(clean) < 2:
                return None
            search_url = f"https://en.wikipedia.org/w/api.php?action=opensearch&search={urllib.parse.quote(clean)}&limit=1&namespace=0&format=json"
            req = urllib.request.Request(search_url, headers={"User-Agent": "Gap2GrowAcademic/1.0 (academic; student@vignan.ac.in)"})
            res = json.loads(urllib.request.urlopen(req, timeout=3.5).read().decode("utf-8"))
            if not res or len(res) < 2 or not res[1]:
                return None
            title = res[1][0]
            extract_url = f"https://en.wikipedia.org/w/api.php?action=query&prop=extracts&exintro&explaintext&titles={urllib.parse.quote(title)}&format=json"
            req2 = urllib.request.Request(extract_url, headers={"User-Agent": "Gap2GrowAcademic/1.0 (academic; student@vignan.ac.in)"})
            data = json.loads(urllib.request.urlopen(req2, timeout=3.5).read().decode("utf-8"))
            pages = data.get("query", {}).get("pages", {})
            for _, page in pages.items():
                extract = page.get("extract", "")
                if extract and len(extract) > 40:
                    return {"title": title, "summary": extract}
        except Exception:
            return None
        return None

    def _call_llm(self, message: str, user_context: dict) -> Optional[str]:
        system_instruction = f"You are the Gap2Grow AI Learning Assistant at Vignan University for {user_context.get('name', 'Student')} ({user_context.get('department', 'CSE')}). Answer clearly with code, tables, and step-by-step explanations where relevant."
        
        # 1. Check Groq
        groq_key = user_context.get("apiKey") or self.groq_key or os.getenv("GROQ_API_KEY")
        if groq_key:
            models = ['groq/compound', 'qwen/qwen3.8-27b', 'groq/compound-mini', 'openai/gpt-oss-120b', 'llama-3.3-70b-versatile', 'llama3-70b-8192']
            if GROQ_AVAILABLE:
                try:
                    client = Groq(api_key=groq_key)
                    for m in models:
                        try:
                            comp = client.chat.completions.create(
                                model=m,
                                messages=[
                                    {"role": "system", "content": system_instruction},
                                    {"role": "user", "content": message}
                                ],
                                temperature=0.3,
                                max_tokens=1500
                            )
                            if comp and comp.choices and comp.choices[0].message.content:
                                return comp.choices[0].message.content
                        except Exception:
                            continue
                except Exception:
                    pass

            # Groq REST Fallback
            url = "https://api.groq.com/openai/v1/chat/completions"
            headers = {"Content-Type": "application/json", "Authorization": f"Bearer {groq_key}"}
            for m in models:
                try:
                    body = {
                        "model": m,
                        "messages": [
                            {"role": "system", "content": system_instruction},
                            {"role": "user", "content": message}
                        ],
                        "temperature": 0.3,
                        "max_tokens": 1500
                    }
                    req = urllib.request.Request(url, data=json.dumps(body).encode("utf-8"), headers=headers, method="POST")
                    with urllib.request.urlopen(req, timeout=12) as resp:
                        res_json = json.loads(resp.read().decode("utf-8"))
                        text = res_json["choices"][0]["message"]["content"]
                        if text:
                            return text
                except Exception:
                    continue

        # 2. Check Gemini
        gemini_key = user_context.get("apiKey") or self.gemini_key or os.getenv("GEMINI_API_KEY")
        if gemini_key and not gemini_key.startswith("gsk_") and GENAI_AVAILABLE:
            try:
                client = genai.Client(api_key=gemini_key)
                prompt = f"{system_instruction}\n\nStudent Query: {message}"
                for m_name in ["gemini-3.6-flash", "gemini-3.5-flash", "gemini-2.5-flash", "gemini-2.0-flash", "gemini-1.5-flash"]:
                    try:
                        response = client.models.generate_content(
                            model=m_name,
                            contents=prompt
                        )
                        if response and response.text:
                            return response.text
                    except Exception:
                        continue
            except Exception:
                pass

        # 3. Check OpenAI
        openai_key = self.openai_key or os.getenv("OPENAI_API_KEY")
        if openai_key:
            try:
                url = "https://api.openai.com/v1/chat/completions"
                payload = json.dumps({
                    "model": "gpt-4o-mini",
                    "messages": [
                        {"role": "system", "content": system_instruction},
                        {"role": "user", "content": message}
                    ]
                }).encode("utf-8")
                req = urllib.request.Request(url, data=payload, headers={
                    "Content-Type": "application/json",
                    "Authorization": f"Bearer {openai_key}"
                })
                data = json.loads(urllib.request.urlopen(req, timeout=8).read().decode("utf-8"))
                return data["choices"][0]["message"]["content"]
            except Exception:
                pass

        return None

    def chat_assistant(self, message: str, user_context: dict, session_id: str = "default") -> dict:
        """
        Delegates all student queries to academic_ai_agent for robust intent classification,
        primary/backup provider fallback, database resource lookup, response validation,
        and multi-turn chat memory context.
        """
        from ai_agent import academic_ai_agent
        result = academic_ai_agent.process_query(message, user_context, session_id=session_id)
        return {
            "reply": result.get("reply", ""),
            "intent": result.get("intent", "GENERAL_STUDENT_QUERY"),
            "mode": result.get("mode", "AI_AGENT"),
            "language": result.get("language", "ENGLISH"),
            "sender": "Gap2Grow AI Learning Assistant",
            "timestamp": "Just now"
        }

    # =========================================================================
    # SUBJECT-WISE TOPIC LOGIC & YOUTUBE RELEVANCE VALIDATION LAYER
    # =========================================================================

    def _get_subject_key(self, subject_input: Optional[str]) -> str:
        if not subject_input:
            return "dbms"
        s = subject_input.lower().strip()
        if "data structure" in s or "ds" == s or s == "data-structures":
            return "data-structures"
        elif "dbms" in s or "database" in s:
            return "dbms"
        elif "object oriented" in s or "oop" in s or "java" in s:
            return "oop"
        elif "artificial" in s or "ai" in s or "intelligence" in s:
            return "ai"
        elif "discrete" in s or "math" in s or s == "discrete-math":
            return "discrete-math"
        elif "digital logic" in s or "dld" in s or s == "digital-logic":
            return "digital-logic"
        elif "visualization" in s or "wrangling" in s or "data viz" in s or s == "data-visualization":
            return "data-visualization"
        elif "operating" in s or "os" == s:
            return "operating-systems"
        elif "network" in s or "cn" == s:
            return "computer-networks"
        return "dbms"

    def validate_video_relevance(self, subject_key: str, topic_name: str, video_title: str, video_desc: str) -> bool:
        """
        Relevance validation layer:
        Input: subject, topic, video title, video description.
        Checks whether video is actually related to the selected subject and topic.
        Rejects videos containing unrelated concepts (cross-subject leakage).
        """
        title_lower = (video_title or "").lower()
        desc_lower = (video_desc or "").lower()
        text = f"{title_lower} {desc_lower}"

        forbidden_map = {
            "dbms": ["array", "linked list", "binary tree", "stack", "queue", "cpu scheduling", "deadlock", "operating system"],
            "data-structures": ["relational algebra", "sql query", "1nf 2nf 3nf", "cpu scheduling", "deadlock", "gate level logic"],
            "oop": ["relational algebra", "sql join", "deadlock", "k-map", "minterm", "osi model", "tcp/ip"],
            "operating-systems": ["relational algebra", "sql join", "pandas dataframe", "1nf 2nf 3nf", "linked list node"],
            "computer-networks": ["relational algebra", "sql join", "binary tree traversal", "cpu scheduling", "k-map"],
            "digital-logic": ["relational algebra", "sql query", "binary tree traversal", "osi model", "pandas dataframe"],
            "data-visualization": ["relational algebra", "sql query", "cpu scheduling", "k-map", "deadlock"]
        }

        # Check forbidden keywords for the selected subject
        for forbidden in forbidden_map.get(subject_key, []):
            if forbidden in text:
                return False

        # Positive keywords check
        topic_words = [w for w in re.findall(r'\b[a-z0-9]+\b', topic_name.lower()) if len(w) > 2 and w not in ["and", "for", "the", "with"]]
        matches = [w for w in topic_words if w in text]
        if len(topic_words) > 0 and len(matches) == 0:
            # If none of the topic words match, check if subject words match
            subject_words = [w for w in re.findall(r'\b[a-z0-9]+\b', subject_key.replace('-', ' ')) if len(w) > 2]
            subj_matches = [w for w in subject_words if w in text]
            if len(subj_matches) == 0:
                return False

        return True

    def search_youtube_videos(self, subject: str, topic: str, subtopic: Optional[str] = None) -> dict:
        """
        Backend YouTube search service.
        Query constructed as: "${subject} ${topic} tutorial" or "${subject} ${topic} ${subtopic} tutorial"
        Validates relevance and rejects cross-subject unrelated videos.
        """
        subject_key = self._get_subject_key(subject)
        sub_data = SUBJECTS_DATA.get(subject_key, SUBJECTS_DATA["dbms"])

        search_query = f"{sub_data['name']} {topic} {subtopic or ''} tutorial".strip()
        encoded_query = urllib.parse.quote(search_query)
        encoded_subject = urllib.parse.quote(sub_data['name'])
        encoded_topic = urllib.parse.quote(topic)
        fallback_search_url = f"https://www.youtube.com/results?search_query={encoded_subject}+{encoded_topic}"

        # 1. Try YouTube Data API v3 if API key available
        yt_api_key = os.getenv("YOUTUBE_API_KEY", "AQ.Ab8RN6LqSEQF2-O5HbFu45bDyU8pPJB_HrcOG45aItU1SYkNGw")
        if yt_api_key:
            try:
                yt_url = f"https://www.googleapis.com/youtube/v3/search?part=snippet&maxResults=5&q={encoded_query}&type=video&key={yt_api_key}"
                req = urllib.request.Request(yt_url, headers={"User-Agent": "Gap2Grow/1.0"})
                with urllib.request.urlopen(req, timeout=5) as resp:
                    res_data = json.loads(resp.read().decode("utf-8"))
                    items = res_data.get("items", [])
                    validated_videos = []
                    for item in items:
                        v_id = item.get("id", {}).get("videoId")
                        snippet = item.get("snippet", {})
                        v_title = snippet.get("title", "")
                        v_desc = snippet.get("description", "")
                        if v_id and self.validate_video_relevance(subject_key, topic, v_title, v_desc):
                            validated_videos.append({
                                "videoId": v_id,
                                "title": v_title,
                                "channelTitle": snippet.get("channelTitle", "Verified Educational Channel"),
                                "thumbnail": snippet.get("thumbnails", {}).get("high", {}).get("url", f"https://img.youtube.com/vi/{v_id}/hqdefault.jpg"),
                                "description": v_desc,
                                "url": f"https://www.youtube.com/watch?v={v_id}"
                            })

                    if validated_videos:
                        return {"videos": validated_videos, "searchUrl": fallback_search_url}
            except Exception as e:
                print(f"YouTube API call failed: {e}")

        # 2. Check topic-specific curated videos in SUBJECTS_DATA
        for t_obj in sub_data["topics"]:
            if topic.lower() in t_obj["name"].lower() or t_obj["name"].lower() in topic.lower():
                for sub_obj in t_obj.get("subtopics", []):
                    if not subtopic or subtopic.lower() in sub_obj["name"].lower():
                        v_info = sub_obj.get("video", {})
                        if v_info and v_info.get("youtubeUrl"):
                            v_url = v_info["youtubeUrl"]
                            v_id = self.extract_youtube_id(v_url) or "video_fallback"
                            return {
                                "videos": [
                                    {
                                        "videoId": v_id,
                                        "title": v_info.get("title", f"{sub_data['name']} - {topic}"),
                                        "channelTitle": v_info.get("channel", "Verified Gate / College Instructor"),
                                        "thumbnail": f"https://img.youtube.com/vi/{v_id}/hqdefault.jpg",
                                        "description": f"Verified tutorial covering {topic} in {sub_data['name']}",
                                        "url": f"https://www.youtube.com/watch?v={v_id}" if v_id != "video_fallback" else v_url
                                    }
                                ],
                                "searchUrl": fallback_search_url
                            }

        # 3. If no relevant video passes validation, return noVideoFound indicator
        return {
            "videos": [],
            "noVideoFound": True,
            "message": "No topic-specific video found yet.",
            "searchUrl": fallback_search_url
        }

    def get_topic_gaps(self, subject: Optional[str] = None) -> list:
        """Returns subject-wise topic list (min 5 topics per subject)."""
        subject_key = self._get_subject_key(subject)
        sub_data = SUBJECTS_DATA.get(subject_key, SUBJECTS_DATA["dbms"])
        
        # Realistic default skill level calibration (Strong, Medium, Weak, Critical)
        skill_distribution = ["Strong", "Weak", "Medium", "Medium", "Critical"]
        progress_distribution = [88, 35, 65, 60, 25]

        result = []
        for idx, t in enumerate(sub_data.get("topics", [])):
            desc = t.get("shortDescription") or t.get("description", "")
            skill = skill_distribution[idx % len(skill_distribution)]
            pct = progress_distribution[idx % len(progress_distribution)]
            gap = "Critical" if skill == "Critical" else ("High" if skill == "Weak" else ("Medium" if skill == "Medium" else "Low"))

            result.append({
                "id": t["id"],
                "name": t["name"],
                "shortDescription": desc,
                "description": desc,
                "skillLevel": skill,
                "progressPct": pct,
                "gapLevel": gap,
                "accuracy": pct,
                "isGap": gap in ["High", "Critical"],
                "subtopics": t.get("subtopics", [])
            })
        return result

    def get_before_after_data(self, subject: Optional[str] = None) -> dict:
        """Returns subject-wise before vs after analytics data."""
        if subject and subject.strip():
            subject_key = self._get_subject_key(subject)
            sub_data = SUBJECTS_DATA.get(subject_key, SUBJECTS_DATA["dbms"])
            ba = sub_data.get("beforeAfter", {})
            topic_items = []
            if "topics" in ba and ba["topics"]:
                for t in ba["topics"]:
                    topic_items.append({
                        "topic": t.get("topic", t.get("name")),
                        "before": t.get("before", 40),
                        "after": t.get("after", 80),
                        "gain": t.get("gain", 40)
                    })
            else:
                for t in sub_data.get("topics", []):
                    after_val = t.get("progressPct", 80)
                    before_val = max(20, after_val - 40)
                    topic_items.append({
                        "topic": t.get("name", "Topic"),
                        "before": before_val,
                        "after": after_val,
                        "gain": after_val - before_val
                    })

            b_avg = ba.get("diagnosticScore")
            if b_avg is None and topic_items:
                b_avg = round(sum(t["before"] for t in topic_items) / len(topic_items))
            elif b_avg is None:
                b_avg = 40

            a_avg = ba.get("currentScore")
            if a_avg is None and topic_items:
                a_avg = round(sum(t["after"] for t in topic_items) / len(topic_items))
            elif a_avg is None:
                a_avg = 80

            return {
                "subject": sub_data["name"],
                "summary": {
                    "beforeAvg": b_avg,
                    "afterAvg": a_avg,
                    "overallGain": a_avg - b_avg
                },
                "topics": topic_items
            }

        # All subjects overview
        all_topics = []
        for key, data in SUBJECTS_DATA.items():
            ba = data.get("beforeAfter", {})
            diag = ba.get("diagnosticScore", 40)
            curr = ba.get("currentScore", 80)
            all_topics.append({
                "topic": data["name"],
                "before": diag,
                "after": curr,
                "gain": curr - diag
            })

        avg_before = round(sum(t["before"] for t in all_topics) / max(len(all_topics), 1))
        avg_after = round(sum(t["after"] for t in all_topics) / max(len(all_topics), 1))

        return {
            "subject": "All Subjects",
            "summary": {
                "beforeAvg": avg_before,
                "afterAvg": avg_after,
                "overallGain": avg_after - avg_before
            },
            "topics": all_topics
        }

    def generate_smart_path(self, student_name: Optional[str] = "Rahul Kumar", gap_topic: Optional[str] = None, subject: Optional[str] = None) -> dict:
        """Returns subject-specific ZERO -> HERO learning path and playlist URL."""
        subject_key = self._get_subject_key(subject)
        sub_data = SUBJECTS_DATA.get(subject_key, SUBJECTS_DATA["dbms"])

        return {
            "studentName": student_name or "Rahul Kumar",
            "subject": sub_data["name"],
            "playlistUrl": sub_data.get("playlistUrl", f"https://www.youtube.com/results?search_query={urllib.parse.quote(sub_data['name'] + ' Zero to Hero complete course')}"),
            "stages": sub_data.get("zeroToHeroPath", [])
        }

    def get_practice_questions(self, subject: Optional[str] = None, topic: Optional[str] = None, subtopic: Optional[str] = None) -> list:
        """Returns 10 topic-specific practice questions."""
        subject_key = self._get_subject_key(subject)
        sub_data = SUBJECTS_DATA.get(subject_key, SUBJECTS_DATA.get("dbms", {}))

        target_topic = None
        if topic:
            topic_clean = topic.lower().replace("&", "and").strip()
            # 1. Exact or substring match
            for t in sub_data.get("topics", []):
                t_clean = t["name"].lower().replace("&", "and").strip()
                if topic_clean in t_clean or t_clean in topic_clean:
                    target_topic = t
                    break
            
            # 2. Key word match if no match found
            if not target_topic:
                topic_words = [w for w in re.findall(r'\b[a-z0-9]+\b', topic_clean) if len(w) > 2 and w not in ["and", "for", "the", "with"]]
                for t in sub_data.get("topics", []):
                    t_clean = t["name"].lower().replace("&", "and").strip()
                    if any(w in t_clean for w in topic_words):
                        target_topic = t
                        break

        if not target_topic and sub_data.get("topics"):
            target_topic = sub_data["topics"][0]

        if not target_topic:
            return []

        # Check subtopics if requested specifically
        if subtopic:
            for sub in target_topic.get("subtopics", []):
                if subtopic.lower() in sub["name"].lower():
                    return sub.get("practiceQuestions", [])[:10]

        # Collect questions across target topic level and subtopics
        questions = []
        if "practiceQuestions" in target_topic:
            questions.extend(target_topic.get("practiceQuestions", []))

        for sub in target_topic.get("subtopics", []):
            questions.extend(sub.get("practiceQuestions", []))

        # Deduplicate while preserving order
        seen = set()
        unique_questions = []
        for q in questions:
            q_text = q.get("question", "")
            if q_text and q_text not in seen:
                seen.add(q_text)
                unique_questions.append(q)

        return unique_questions[:10] if unique_questions else []


    def get_topic_hierarchy(self, subject: Optional[str] = None, topic: Optional[str] = None, subtopic: Optional[str] = None, topic_id: Optional[str] = None) -> dict:
        """Returns step-by-step topic hierarchy for selected subject and topic with 4 stages (1 video per stage)."""
        subject_key = self._get_subject_key(subject or topic_id)
        sub_data = SUBJECTS_DATA.get(subject_key, SUBJECTS_DATA["dbms"])

        target_topic = None
        if topic:
            for t in sub_data["topics"]:
                if topic.lower() in t["name"].lower() or t["name"].lower() in topic.lower():
                    target_topic = t
                    break

        if not target_topic:
            target_topic = sub_data["topics"][0]

        subtopics = target_topic.get("subtopics", [])
        desc = target_topic.get("shortDescription") or target_topic.get("description", "")

        hierarchy_stages = []
        for idx in range(4):
            if idx < len(subtopics):
                sub = subtopics[idx]
            else:
                sub = {
                    "name": f"{target_topic['name']} Step {idx + 1}",
                    "simpleExplanation": f"Stage {idx + 1} learning milestone and concepts for {target_topic['name']}."
                }
            hierarchy_stages.append({
                "category": f"Stage {idx + 1}: {sub.get('name', f'Step {idx + 1}')}",
                "description": sub.get("simpleExplanation", f"Learning milestone Stage {idx + 1} for {target_topic['name']}"),
                "subtopics": [sub]
            })

        return {
            "subject": sub_data["name"],
            "topic": target_topic["name"],
            "topicId": topic_id if topic_id else target_topic.get("id", target_topic["name"]),
            "overview": desc,
            "hierarchy": hierarchy_stages
        }


ai_service = AIService()

